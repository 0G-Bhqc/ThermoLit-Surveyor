"""
export.py — 调研报告 Markdown → docx(python-docx,可选依赖 export)。

支持元素:标题(#~####)、段落(**粗体** 内联)、引用块(>)、无序列表、
表格(| 分隔,含表头加粗)。刻意保持转换器简单——报告由 synthesize 节点
按固定章节结构生成,这里只做忠实排版,不重写内容。
"""
from __future__ import annotations

import re
from pathlib import Path


def _add_runs(paragraph, text: str) -> None:
    """按 **粗体** 标记拆分内联 run。"""
    for i, seg in enumerate(text.split("**")):
        if not seg:
            continue
        run = paragraph.add_run(seg)
        if i % 2 == 1:
            run.bold = True


def _is_table_separator(line: str) -> bool:
    stripped = line.strip()
    if not stripped.startswith("|"):
        return False
    cells = [c.strip() for c in stripped.strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def report_to_docx(report_md: str, out_path: str | Path,
                   title: str | None = None) -> Path:
    """把 Markdown 报告转换为 docx。返回输出路径。"""
    from docx import Document

    doc = Document()
    if title:
        doc.add_heading(title, level=0)

    lines = report_md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|") and i + 1 < len(lines) \
                and _is_table_separator(lines[i + 1]):
            header = _split_row(lines[i])
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(_split_row(lines[j]))
                j += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Light Grid Accent 1"
            for c, text in enumerate(header):
                run = table.rows[0].cells[c].paragraphs[0].add_run(
                    re.sub(r"\*\*", "", text))
                run.bold = True
            for r, row in enumerate(rows, 1):
                for c in range(len(header)):
                    table.rows[r].cells[c].text = re.sub(r"\*\*", "", row[c]) \
                        if c < len(row) else ""
            i = j
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            doc.add_heading(stripped.lstrip("#").strip(), level=min(level, 4))
        elif stripped.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            _add_runs(p, stripped.lstrip(">").strip())
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, stripped[2:])
        elif not stripped:
            pass
        else:
            p = doc.add_paragraph()
            _add_runs(p, stripped)
        i += 1

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out
