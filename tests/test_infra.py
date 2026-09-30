"""token 计量代理、docx 导出、断点续跑、FastAPI 服务的离线测试。"""
import json
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from thermolit.graph import UsageTracker, _TrackingLLM


# ---------------------------------------------------------------- 计量代理
class _Msg:
    def __init__(self, tokens):
        self.usage_metadata = {"input_tokens": tokens[0], "output_tokens": tokens[1]}


class _Inner:
    def __init__(self):
        self.calls = 0

    def invoke(self, messages, **kwargs):
        self.calls += 1
        return _Msg((10, 5))


class TestUsageTracker(unittest.TestCase):
    def test_accumulates(self):
        inner = _Inner()
        tracker = UsageTracker()
        llm = _TrackingLLM(inner, tracker)
        llm.invoke([{"x": 1}])
        llm.invoke([{"x": 2}])
        self.assertEqual(inner.calls, 2)
        d = tracker.as_dict()
        self.assertEqual(d["input_tokens"], 20)
        self.assertEqual(d["output_tokens"], 10)
        self.assertEqual(d["total_tokens"], 30)
        self.assertEqual(d["requests"], 2)

    def test_missing_usage_tolerated(self):
        class _NoUsage:
            pass

        class _Inner2:
            def invoke(self, messages, **kwargs):
                m = _NoUsage()
                m.usage_metadata = None
                return m

        tracker = UsageTracker()
        _TrackingLLM(_Inner2(), tracker).invoke([])
        self.assertEqual(tracker.as_dict()["requests"], 1)
        self.assertEqual(tracker.as_dict()["total_tokens"], 0)


# ---------------------------------------------------------------- docx 导出
class TestExportDocx(unittest.TestCase):
    REPORT = """# 测试报告

> 引用块:**数值**以审计表为准。

- 要点一
- 要点二

| A | B |
|---|---|
| 1 | 2 |

正文段落,包含 **粗体** 文本。
"""

    def test_conversion(self):
        from docx import Document

        from thermolit.export import report_to_docx
        out = report_to_docx(self.REPORT, "eval/results/test_export.docx",
                             title="冒烟导出")
        self.assertTrue(out.exists())
        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(any("测试报告" in t for t in texts))
        self.assertTrue(any("要点一" in t for t in texts))
        self.assertTrue(any("粗体" in t for t in texts))
        # 表格:表头 + 1 行数据
        self.assertEqual(len(doc.tables[0].rows), 2)
        self.assertEqual(doc.tables[0].rows[1].cells[0].text, "1")
        out.unlink()


# ---------------------------------------------------------------- 断点续跑
class TestCheckpoint(unittest.TestCase):
    def test_checkpoint_creates_sqlite(self):
        try:
            import langgraph.checkpoint.sqlite  # noqa: F401
        except ImportError:
            self.skipTest("langgraph-checkpoint-sqlite 未安装")

        import contextlib

        from test_graph_offline import FakeAdapter, FakeLLM

        from thermolit import graph as proto

        # pass1 ×1 + pass2 ×2 = 3 份抽取输出(避免重试延时)
        record = {"material_system": "Bi2Te3", "zT_value": "0.9 @ 350 K",
                  "seebeck_coefficient": "-180 uV/K",
                  "electrical_conductivity": "1200 S/cm",
                  "thermal_conductivity": "1.3 W/mK",
                  "temperature_k": 350, "key_claim": "x"}
        fake_llm = FakeLLM([json.dumps(dict(record)) for _ in range(3)])
        db = Path("eval/results") / f"test_checkpoint_{time.time_ns()}.db"
        db.parent.mkdir(parents=True, exist_ok=True)
        try:
            with patch.object(proto, "_make_llm", return_value=fake_llm), \
                 patch.object(proto, "get_adapter", return_value=FakeAdapter()):
                state = proto.run_agent("Bi2Te3 checkpoint smoke", top_k=1,
                                        num_followup=2, check_dois=False,
                                        checkpoint_path=str(db))
            self.assertIn("deep_gap_report", state)
            self.assertTrue(db.exists(), "checkpoint 数据库应已创建")
            self.assertGreater(db.stat().st_size, 0)
            # 同一 checkpoint 第二次运行(同 thread_id)不崩
            with patch.object(proto, "_make_llm", return_value=fake_llm), \
                 patch.object(proto, "get_adapter", return_value=FakeAdapter()):
                state2 = proto.run_agent("Bi2Te3 checkpoint smoke", top_k=1,
                                         num_followup=2, check_dois=False,
                                         checkpoint_path=str(db))
            self.assertIn("metrics", state2)
        finally:
            with contextlib.suppress(OSError):
                db.unlink()  # Windows 下连接未释放时留给 gitignore,不致失败


# ---------------------------------------------------------------- FastAPI 服务
class TestApi(unittest.TestCase):
    def _client(self):
        from fastapi.testclient import TestClient

        from thermolit.api import app
        return TestClient(app)

    def test_health_and_job_flow(self):
        try:
            import fastapi  # noqa: F401
        except ImportError:
            self.skipTest("fastapi 未安装")
        import thermolit.graph as proto

        fake_state = {"material_system": "Bi2Te3", "deep_gap_report": "# 报告",
                      "audit_table_md": "| t |", "hypotheses": [],
                      "experiment_plan": {}, "doi_check": {}, "metrics": {}}
        client = self._client()
        self.assertEqual(client.get("/health").json()["status"], "ok")
        self.assertEqual(client.get("/jobs/nope").status_code, 404)

        with patch.object(proto, "run_agent", return_value=fake_state):
            resp = client.post("/survey", json={"query": "Bi2Te3 offline"})
            self.assertEqual(resp.status_code, 200)
            job_id = resp.json()["job_id"]
            deadline = time.time() + 5
            while time.time() < deadline:
                job = client.get(f"/jobs/{job_id}").json()
                if job["status"] != "running":
                    break
                time.sleep(0.05)
            self.assertEqual(job["status"], "done")
            self.assertEqual(job["report"], "# 报告")
            self.assertIn("metrics", job)


if __name__ == "__main__":
    unittest.main()
